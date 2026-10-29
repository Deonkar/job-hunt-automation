import os
import sys
import tempfile
import unittest

import docx

sys.path.insert(0, os.path.dirname(__file__))
import build_resume as br  # noqa: E402
import make_master as mm  # noqa: E402

# A made-up candidate: tests never read anyone's real resume.
PROFILE = {
    "name": "Alex Example",
    "headline": "Software Engineer",
    "location": "Springfield, Exampleland",
    "email": "alex@example.com",
    "phone": "+00 0000 000000",
    "links": ["https://linkedin.com/in/alex-example", "https://github.com/alex-example"],
    "summary": "Backend engineer with **3 years** shipping payment APIs.",
    "skills": {
        "Languages": ["Go", "Python", "TypeScript", "Java"],
        "Data": ["PostgreSQL", "Redis"],
        "AI / ML": ["RAG", "LLM tool calling"],
    },
    "experience": {
        "full_time": [
            {"key": "acme", "company": "Acme Corp", "title": "Software Engineer", "location": "Springfield",
             "start": "2023-01", "end": None},
            {"key": "initech", "company": "Initech", "title": "Junior Developer", "location": "Remote",
             "start": "2021-06", "end": "2022-12"},
        ],
    },
    "projects": [{"key": "tool", "name": "Toolbox", "stack": "Go, Redis"},
                 {"key": "bot", "name": "Chatbot", "stack": "Python, RAG"}],
    "education": [{"degree": "B.Sc. Computer Science", "school": "Example University", "grade": "GPA 3.6/4",
                   "years": "2017 - 2021"}],
}
BANK = {
    "experience": {
        "acme": {
            "pay": {"base": "Built the **payments API**: 1.2M requests a day.",
                    "short": "Built the **payments API** (1.2M req/day)."},
            "lat": {"base": "Cut p99 latency from 900ms to 120ms with Redis caching."},
            "ai": {"base": "Shipped a RAG support assistant used by 40 agents."},
        },
        "initech": {
            "crud": {"base": "Maintained internal TypeScript admin tools."},
        },
    },
    "projects": {
        "tool": {"tb": {"base": "CLI toolbox with 300 GitHub stars.", "short": "CLI toolbox, 300 stars."}},
        "bot": {"cb": {"base": "Chatbot answering docs questions with RAG."}},
    },
    "internships": [
        {"title": "SWE Intern", "company": "Globex", "dates": "Jun 2020 - Aug 2020"},
        {"title": "Data Intern", "company": "Umbrella", "dates": "Jan 2020 - Mar 2020", "note": "remote"},
    ],
}

_DIR = tempfile.mkdtemp()
MASTER = mm.make(PROFILE, BANK, os.path.join(_DIR, "master.docx"))


def plain(markup):
    return markup.replace("**", "")


def build(patch):
    out = os.path.join(tempfile.mkdtemp(), "out.docx")
    br.build(MASTER, BANK, patch, out, location=PROFILE["location"])
    return docx.Document(out)


def texts(doc):
    return [p.text for p in doc.paragraphs]


class MarkupTest(unittest.TestCase):
    def test_splits_bold_segments(self):
        self.assertEqual(
            br.parse_markup("a **b** c"),
            [("a ", False), ("b", True), (" c", False)],
        )


class MasterTest(unittest.TestCase):
    def test_layout_the_builder_expects(self):
        t = texts(docx.Document(MASTER))
        self.assertEqual(t[0], "Alex Example")
        self.assertEqual(t[1], "Software Engineer")
        self.assertIn("Springfield, Exampleland", t[2])
        self.assertIn("alex@example.com", t[2])
        for h in ("Summary", "Skills", "Experience", "Projects", "Education"):
            self.assertIn(h, t)
        self.assertIn("Languages: Go, Python, TypeScript, Java", t)
        self.assertTrue(any(x.endswith("\tJan 2023 - Present") for x in t))

    def test_every_bank_bullet_is_a_numbered_paragraph(self):
        doc = docx.Document(MASTER)
        numbered = [p.text for p in doc.paragraphs if br._is_numbered(p)]
        self.assertEqual(len(numbered), 6)
        self.assertIn(plain(BANK["experience"]["acme"]["pay"]["base"]), numbered)

    def test_contact_links_are_real_hyperlinks(self):
        doc = docx.Document(MASTER)
        urls = [r.target_ref for r in doc.part.rels.values() if r.is_external]
        self.assertEqual(sorted(urls), sorted(PROFILE["links"]))

    def test_no_projects_section_when_none(self):
        p = dict(PROFILE, projects=[])
        b = dict(BANK, projects={})
        t = texts(docx.Document(mm.make(p, b, os.path.join(tempfile.mkdtemp(), "m.docx"))))
        self.assertNotIn("Projects", t)


class BuildTest(unittest.TestCase):
    def test_empty_patch_keeps_master_text_except_contact_links(self):
        master = texts(docx.Document(MASTER))
        out = texts(build({}))
        self.assertEqual(len(master), len(out))
        diffs = [i for i, (a, b) in enumerate(zip(master, out)) if a != b]
        self.assertEqual(diffs, [3])
        self.assertIn("linkedin.com/in/alex-example", out[3])
        self.assertNotIn("https://", out[3])

    def test_applies_title_summary_and_per_role_bullets(self):
        doc = build({
            "title": "Backend Software Engineer",
            "summary": "Engineer with **big numbers** here.",
            "experience": {"acme": ["ai", "pay.short"]},
        })
        t = texts(doc)
        self.assertEqual(t[1], "Backend Software Engineer")
        summary = next(p for p in doc.paragraphs if p.text.startswith("Engineer with"))
        self.assertTrue(any(r.bold and r.text == "big numbers" for r in summary.runs))
        blocks = br.blocks(doc, "Experience", ["acme", "initech"])
        acme = [p.text for p in blocks["acme"] if br._is_numbered(p)]
        self.assertEqual(acme, [plain(BANK["experience"]["acme"]["ai"]["base"]),
                                plain(BANK["experience"]["acme"]["pay"]["short"])])
        initech = [p.text for p in blocks["initech"] if br._is_numbered(p)]
        self.assertEqual(initech, [plain(BANK["experience"]["initech"]["crud"]["base"])])

    def test_reorders_projects_and_picks_their_bullets(self):
        t = texts(build({"projects_order": ["bot", "tool"], "projects": {"tool": ["tb.short"]}}))
        bot = next(i for i, x in enumerate(t) if x.startswith("Chatbot"))
        tool = next(i for i, x in enumerate(t) if x.startswith("Toolbox"))
        self.assertLess(bot, tool)
        self.assertEqual(t[tool + 1], plain(BANK["projects"]["tool"]["tb"]["short"]))

    def test_internships_section_sits_before_projects(self):
        t = texts(build({"internships": True}))
        i = t.index("Internships")
        self.assertLess(i, t.index("Projects"))
        self.assertEqual(t[i + 1], "SWE Intern, Globex\tJun 2020 - Aug 2020")
        self.assertEqual(t[i + 2], "Data Intern, Umbrella - remote\tJan 2020 - Mar 2020")

    def test_reorders_skill_lines_and_items(self):
        t = texts(build({
            "skills_order": ["AI / ML", "Languages"],
            "skills_items": {"Languages": ["Python", "Go", "Java", "TypeScript"]},
        }))
        s = t.index("Skills")
        self.assertTrue(t[s + 1].startswith("AI / ML: "))
        self.assertEqual(t[s + 2], "Languages: Python, Go, Java, TypeScript")
        self.assertEqual(sum(x.startswith("Data: ") for x in t), 1)

    def test_rejects_unknown_bullet_role_and_key(self):
        for bad in ({"experience": {"acme": ["nope"]}}, {"experience": {"nobody": ["pay"]}}, {"old_key": ["pay"]}):
            with self.assertRaises(ValueError, msg=bad):
                build(bad)

    def test_rejects_skill_not_on_master(self):
        with self.assertRaises(ValueError):
            build({"skills_items": {"Languages": ["Rust", "Go", "Python", "TypeScript", "Java"]}})

    def test_rejects_dropping_a_skill(self):
        with self.assertRaises(ValueError):
            build({"skills_items": {"Languages": ["Go", "Python"]}})


def build_country(country, patch=None, spoken=None):
    out = os.path.join(tempfile.mkdtemp(), "out.docx")
    br.build(MASTER, BANK, patch or {}, out, country=country, spoken=spoken, location=PROFILE["location"])
    return docx.Document(out)


class CountryTest(unittest.TestCase):
    def test_sets_a4_page_keeping_text_width(self):
        doc = build_country({"page": "A4"})
        self.assertAlmostEqual(doc.sections[0].page_width.mm, 210, delta=0.5)
        self.assertAlmostEqual(doc.sections[0].page_height.mm, 297, delta=0.5)
        s, m = doc.sections[0], docx.Document(MASTER).sections[0]
        width = lambda sec: docx.shared.Emu(sec.page_width - sec.left_margin - sec.right_margin).mm  # noqa: E731
        self.assertAlmostEqual(width(s), width(m), delta=0.1)

    def test_replacements_apply_everywhere(self):
        t = "\n".join(texts(build_country({"replace": [["latency", "response time"]]})))
        self.assertNotIn("latency", t)
        self.assertIn("p99 response time", t)

    def test_replacement_split_across_runs_fails_loudly(self):
        with self.assertRaises(ValueError):
            build_country({"replace": [["payments API: 1.2M", "x"]]})

    def test_regex_dates(self):
        t = texts(build_country({"replace_regex": br.MM_YYYY}, {"internships": True}))
        self.assertTrue(any(x.endswith("\t01/2023 - Present") for x in t))
        self.assertTrue(any(x.endswith("\t06/2020 - 08/2020") for x in t))

    def test_location_suffix_and_work_auth_line(self):
        t = texts(build_country({"location_suffix": " (open to relocation)",
                                 "work_auth": "Work permit: EU Blue Card"}))
        self.assertIn("Springfield, Exampleland (open to relocation)", t[2])
        self.assertEqual(t[4], "Work permit: EU Blue Card")
        self.assertEqual(sum("(open to relocation)" in x for x in t), 1)

    def test_spoken_languages_line_follows_skills(self):
        t = texts(build_country({"spoken_languages": True}, spoken=["English (C1)", "Spanish (native)"]))
        e = t.index("Experience")
        self.assertEqual(t[e - 1], "Spoken languages: English (C1), Spanish (native)")

    def test_output_has_no_em_en_dashes_or_arrows(self):
        t = "\n".join(texts(build_country({}, {"internships": True, "summary": "API 6.7s→450ms — fast"})))
        for ch in ("–", "—", "→"):
            self.assertNotIn(ch, t)
        self.assertIn("API 6.7s to 450ms - fast", t)

    def test_closing_line(self):
        t = texts(build_country({"closing_line": "Referees available on request"}))
        self.assertEqual(t[-1], "Referees available on request")


if __name__ == "__main__":
    unittest.main()
