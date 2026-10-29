"""Render data/master.docx from data/profile.json and data/bullets.json.

The master is the one resume every tailored copy is cloned from, laid out the way
build_resume.py expects: name, headline, contact line (location as its own run), links,
then bold section headings Summary / Skills / Experience / Projects / Education, each role
and project followed by its numbered bullets. Re-run it whenever profile or bullets change.

Usage: py -3 make_master.py [--profile data/profile.json] [--bank data/bullets.json] [--out data/master.docx]
"""
import argparse
import json
import os

import docx
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def month(ym):
    if not ym:
        return "Present"
    y, m = ym.split("-")
    return f"{MONTHS[int(m) - 1]} {y}"


def _markup(p, text, size=None):
    for i, part in enumerate(text.split("**")):
        if part:
            r = p.add_run(part)
            r.bold = i % 2 == 1
            if size:
                r.font.size = Pt(size)
    return p


def _hyperlink(p, url):
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r, t = OxmlElement("w:r"), OxmlElement("w:t")
    t.text = url
    r.append(t)
    h.append(r)
    p._p.append(h)


def make(profile, bank, out):
    doc = docx.Document()
    s = doc.sections[0]
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(0.5)
    right = s.page_width - s.left_margin - s.right_margin
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    bullet_num = doc.styles["List Bullet"].element.pPr.numPr.numId.val

    def para(text="", size=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        return _markup(p, text, size)

    def heading(name):
        p = para(f"**{name}**", 12)
        p.paragraph_format.space_before = Pt(8)

    def tabbed(left, right_text):
        p = para(left)
        p.paragraph_format.tab_stops.add_tab_stop(right, WD_TAB_ALIGNMENT.RIGHT)
        p.add_run("\t")
        p.add_run(right_text)

    def bullets(entries):
        for e in entries.values():
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1)
            num = p._p.get_or_add_pPr().get_or_add_numPr()
            num.get_or_add_ilvl().val = 0
            num.get_or_add_numId().val = bullet_num
            _markup(p, e["base"])

    para(f"**{profile['name']}**", 18)
    para(profile.get("headline", ""))
    contact = para()
    contact.add_run(profile["location"])  # own run: country presets append to it
    for extra in (profile.get("email"), profile.get("phone")):
        if extra:
            contact.add_run(f" | {extra}")
    links = para()
    for i, url in enumerate(profile.get("links", [])):
        if i:
            links.add_run(" | ")
        _hyperlink(links, url)

    heading("Summary")
    para(profile.get("summary", ""))
    heading("Skills")
    for label, items in profile["skills"].items():
        para(f"**{label}: **" + ", ".join(items))
    heading("Experience")
    for role in profile["experience"]["full_time"]:
        para(f"**{role['title']}**")
        tabbed(", ".join(x for x in (role["company"], role.get("location")) if x),
               f"{month(role['start'])} - {month(role.get('end'))}")
        bullets(bank["experience"][role["key"]])
    if profile.get("projects"):
        heading("Projects")
        for proj in profile["projects"]:
            p = para(f"**{proj['name']}**")  # first run = project name
            if proj.get("stack"):
                p.add_run(f" | {proj['stack']}")
            bullets(bank["projects"][proj["key"]])
    heading("Education")
    for e in profile.get("education", []):
        left = f"**{e['degree']}**, {e['school']}" + (f", {e['grade']}" if e.get("grade") else "")
        tabbed(left, e.get("years", ""))
    doc.save(out)
    return out


def main():
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    ap = argparse.ArgumentParser()
    ap.add_argument("--profile", default=os.path.join(root, "data", "profile.json"))
    ap.add_argument("--bank", default=os.path.join(root, "data", "bullets.json"))
    ap.add_argument("--out", default=os.path.join(root, "data", "master.docx"))
    a = ap.parse_args()
    with open(a.profile, encoding="utf-8") as f, open(a.bank, encoding="utf-8") as g:
        print(make(json.load(f), json.load(g), a.out))


if __name__ == "__main__":
    main()
